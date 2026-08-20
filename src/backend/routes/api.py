import os
import json
import tempfile
import logging
import traceback
import base64
import requests
from datetime import datetime
from flask import Blueprint, request, jsonify, current_app
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename

from backend.utils.helpers import allowed_file
from backend.services.story_service import convert_stories_to_frontend_format
from backend.models import User, db
from autoAgile.utils.prompts import (
    extract_text_from_docx, refine_doc, extract_functionarity,
    extract_epics, get_epics, generate_test_cases, refine_requirements, rat,
    extract_epics_v2, get_epics_v2, generate_test_cases_v2,
)
from autoAgile.utils.llm_factory import get_chat_model
from autoAgile.save_output import save_json_output

api_bp = Blueprint('api', __name__)
logger = logging.getLogger(__name__)

# LLM Configuration
LLM_PROVIDER = os.environ.get('LLM_PROVIDER', 'ollama').lower()
GROQ_API_KEY = os.environ.get('GROQ_API_KEY', '').strip()
OPENAI_API_KEY = os.environ.get('OPENAI_API_KEY', '') or os.environ.get('auth_key', '')
OPENAI_API_KEY = OPENAI_API_KEY.strip() if OPENAI_API_KEY else ''
OLLAMA_BASE_URL = os.environ.get('OLLAMA_BASE_URL', 'http://localhost:11434')
OLLAMA_MODEL = os.environ.get('OLLAMA_MODEL', 'llama3.2:latest')

# Use improved v2 prompts by default (set USE_V2_PROMPTS=false in .env to use old prompts)
USE_V2_PROMPTS = os.environ.get('USE_V2_PROMPTS', 'true').lower() == 'true'

@api_bp.route('/health', methods=['GET'])
def health_check():
    """Health check endpoint"""
    return jsonify({
        'status': 'ok', 
        'provider': LLM_PROVIDER,
        'api_key_configured': bool(OPENAI_API_KEY if LLM_PROVIDER == 'openai' else GROQ_API_KEY if LLM_PROVIDER == 'groq' else True),
        'ollama_url': OLLAMA_BASE_URL if LLM_PROVIDER == 'ollama' else None,
        'model': OLLAMA_MODEL if LLM_PROVIDER == 'ollama' else 'llama-3.3-70b-versatile' if LLM_PROVIDER == 'groq' else 'gpt-4-turbo'
    })

@api_bp.route('/generate-stories', methods=['POST'])
@login_required
def generate_stories():
    """Generate user stories from uploaded document (requires authentication)"""
    logger.info("API CALL RECEIVED: /api/generate-stories")
    
    try:
        # Check if file is present
        if 'file' not in request.files:
            return jsonify({'error': 'No file provided'}), 400
        
        file = request.files['file']
        
        if file.filename == '':
            return jsonify({'error': 'No file selected'}), 400
        
        if not allowed_file(file.filename):
            return jsonify({'error': 'File type not allowed. Please upload .docx, .doc, .txt, or .md files'}), 400

        # Validate API key based on provider
        if LLM_PROVIDER == 'openai' and not OPENAI_API_KEY:
            return jsonify({'error': 'OpenAI API key not configured. Set OPENAI_API_KEY in .env'}), 500
        elif LLM_PROVIDER == 'groq' and not GROQ_API_KEY:
            return jsonify({'error': 'Groq API key not configured. Set GROQ_API_KEY in .env'}), 500
        
        # Save uploaded file temporarily
        filename = secure_filename(file.filename)
        
        # Handle macOS resource fork files (._filename) - skip them
        if filename.startswith('._'):
            return jsonify({'error': 'Invalid file: macOS resource fork file detected. Please upload the actual document file, not the resource fork.'}), 400
        
        upload_folder = current_app.config.get('UPLOAD_FOLDER', tempfile.gettempdir())
        filepath = os.path.join(upload_folder, filename)
        
        # Save the file
        file.save(filepath)
        
        # Verify file was saved and exists
        if not os.path.exists(filepath):
            return jsonify({'error': f'Failed to save uploaded file to {filepath}'}), 500
        
        # Verify file is not empty
        if os.path.getsize(filepath) == 0:
            if os.path.exists(filepath):
                os.remove(filepath)
            return jsonify({'error': 'Uploaded file is empty'}), 400
        
        logger.info(f"File saved successfully: {filepath} (size: {os.path.getsize(filepath)} bytes)")
        
        try:
            # Initialize LLM model based on configured provider
            # Higher temperature for more creative/useful stories (0.5-0.7 range)
            # Lower temperature for more deterministic output (0.3-0.4 range)
            temp = float(os.environ.get('LLM_TEMPERATURE', '0.5'))
            
            # Get model name from request or use defaults
            if LLM_PROVIDER == 'ollama':
                model_name = request.form.get('model', OLLAMA_MODEL)
                logger.info(f"[Ollama] Initializing with model: {model_name}")
            elif LLM_PROVIDER == 'openai':
                model_name = request.form.get('model', 'gpt-4-turbo')
                logger.info(f"[OpenAI] Initializing with model: {model_name}")
            elif LLM_PROVIDER == 'groq':
                model_name = "llama-3.3-70b-versatile"
                logger.info("[Groq] Initializing ChatGroq instance...")
            else:
                return jsonify({'error': f'Unknown LLM_PROVIDER: {LLM_PROVIDER}'}), 500
            
            # Use factory to create chat model
            try:
                chat = get_chat_model(temperature=temp, model_name=model_name)
            except ValueError as e:
                return jsonify({'error': str(e)}), 500
            
            if chat is None:
                return jsonify({'error': 'Failed to initialize LLM'}), 500
            
            # Extract text from document - verify file still exists
            if not os.path.exists(filepath):
                raise Exception(f"File was deleted before extraction: {filepath}")
            
            if filename.endswith('.docx') or filename.endswith('.doc'):
                extracted_text = extract_text_from_docx(filepath)
            else:
                with open(filepath, 'r', encoding='utf-8') as f:
                    extracted_text = f.read()
            
            # Process document using autoAgile
            mode = "prod"
            logger.info(f"Processing document: {filename}, length: {len(extracted_text)}")
            
            try:
                logger.info("STEP 1: Extracting requirements...")
                requirements = rat(refine_doc, extract_functionarity, extracted_text, chat, mode)

                if requirements is None:
                    raise Exception("Requirements extraction returned None")

                if not isinstance(requirements, str):
                    requirements = str(requirements)

                # Select which version of functions to use
                extract_epics_func = extract_epics_v2 if USE_V2_PROMPTS else extract_epics
                get_epics_func = get_epics_v2 if USE_V2_PROMPTS else get_epics
                generate_test_cases_func = generate_test_cases_v2 if USE_V2_PROMPTS else generate_test_cases

                # Split requirements into individual lines for batched processing
                BATCH_SIZE = 4
                req_lines = [r.strip() for r in requirements.split('\n')
                             if r.strip() and len(r.strip()) > 15]
                batches = [req_lines[i:i+BATCH_SIZE] for i in range(0, len(req_lines), BATCH_SIZE)] if req_lines else [[requirements]]
                logger.info(f"Step 2: Extracting epics in {len(batches)} batch(es) of up to {BATCH_SIZE} requirements (using {'v2' if USE_V2_PROMPTS else 'v1'} prompts)")

                all_epics = []
                for batch_idx, batch in enumerate(batches):
                    batch_text = '\n'.join(batch)
                    logger.info(f"Step 2 batch {batch_idx + 1}/{len(batches)}: {len(batch)} requirements")
                    batch_result = extract_epics_func(batch_text, chat, mode)
                    if batch_result:
                        try:
                            parsed = json.loads(batch_result)
                            all_epics.extend(parsed.get('Epics', []))
                        except (json.JSONDecodeError, ValueError) as e:
                            logger.warning(f"Batch {batch_idx + 1} epic parse failed: {e}")

                deliverables = json.dumps({'Epics': all_epics}, indent=2)
                logger.info(f"Step 2 complete: {len(all_epics)} stories across all batches")

                if USE_V2_PROMPTS:
                    logger.info("Step 3: Refining epics with detailed definition of done (v2)...")
                    try:
                        epics = get_epics_func(deliverables, chat)
                        if epics is None:
                            raise Exception("Epics refinement returned None")
                        if not isinstance(epics, str):
                            epics = str(epics)
                    except Exception as e:
                        logger.warning(f"Refinement failed, using deliverables directly: {e}")
                        epics = deliverables
                else:
                    logger.info("Step 3: Using deliverables as epics (v1 - refinement skipped)...")
                    epics = deliverables

                logger.info("Step 4: Epics extracted successfully")

                logger.info(f"Step 5: Generating test cases in {len(batches)} batch(es)... (using {'v2' if USE_V2_PROMPTS else 'v1'} prompts)")
                all_tc_groups = []
                tc_counter = 0
                for batch_idx, batch in enumerate(batches):
                    batch_text = '\n'.join(batch)
                    logger.info(f"Step 5 batch {batch_idx + 1}/{len(batches)}: {len(batch)} requirements")
                    batch_tc = generate_test_cases_func(batch_text, chat, mode)
                    if batch_tc:
                        try:
                            parsed = json.loads(batch_tc)
                            groups = parsed.get('test_cases', [])
                            # Renumber TCs globally across batches
                            for group in groups:
                                for tc in group.get('test_cases', []):
                                    tc_counter += 1
                                    tc['id'] = f"TC{tc_counter}"
                            all_tc_groups.extend(groups)
                        except (json.JSONDecodeError, ValueError) as e:
                            logger.warning(f"Batch {batch_idx + 1} TC parse failed: {e}")

                test_cases = json.dumps({'test_cases': all_tc_groups}, indent=2)
                logger.info(f"Step 5 complete: {tc_counter} test cases across all batches")
                    
            
            except Exception as processing_error:
                logger.error(f"Error during processing: {processing_error}")
                logger.error(traceback.format_exc())
                raise
            
            # VALIDATION - Disabled (validation module removed with autoAgile core engine)
            # epics_validation = validate_output(epics, extracted_text)
            # print_validation_report(epics_validation, "Epics/User Stories Validation")
            
            # test_cases_validation = validate_output(test_cases, extracted_text)
            # print_validation_report(test_cases_validation, "Test Cases Validation")
            
            # completeness_validation = validate_requirements_completeness(requirements, epics)
            # print_validation_report(completeness_validation, "Completeness Validation")
            
            # Log epics before conversion for debugging
            logger.info(f"[API] Epics length: {len(epics) if epics else 0}")
            logger.info(f"[API] Epics preview (first 500 chars): {str(epics)[:500] if epics else 'None'}")
            
            # Convert to frontend format
            stories = convert_stories_to_frontend_format(epics, test_cases, requirements)
            
            logger.info(f"[API] Converted stories count: {len(stories) if stories else 0}")
            
            if not stories or len(stories) == 0:
                logger.error("[API] WARNING: No stories generated after conversion!")
                logger.error(f"[API] Epics was: {epics[:1000] if epics else 'None'}")
                # Return error instead of empty success
                return jsonify({
                    'success': False,
                    'error': 'No user stories were generated. Please check the server logs for details.',
                    'stories': [],
                    'count': 0
                }), 500
            
            # Save output to JSON file
            output_file_path = None
            try:
                logger.info(f"[API] Attempting to save output. Requirements type: {type(requirements)}, Epics type: {type(epics)}, Test cases type: {type(test_cases)}")
                logger.info(f"[API] Filepath: {filepath}")
                output_file_path = save_json_output(requirements, epics, test_cases, filepath)
                logger.info(f"Output saved successfully to: {output_file_path}")
            except Exception as save_error:
                logger.error(f"Could not save JSON output: {save_error}")
                logger.error(f"Error traceback: {traceback.format_exc()}")
            
            
            return jsonify({
                'success': True,
                'stories': stories,
                'count': len(stories),
                'output_file': output_file_path
            })
            
        finally:
            # Clean up temporary file
            if os.path.exists(filepath):
                os.remove(filepath)
                
    except Exception as e:
        logger.error(f"Error generating stories: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error generating stories: {str(e)}',
            'details': traceback.format_exc().split('\n')[-5:]
        }), 500

@api_bp.route('/integrate-story', methods=['POST'])
@login_required
def integrate_story():
    """Integrate a single user story - save to json_output (requires authentication)"""
    try:
        data = request.json
        story_id = data.get('storyId')
        story_data = data.get('story')
        
        if not story_data:
            return jsonify({'error': 'Story data not provided'}), 400
        
        # Save integrated story to data/outputs
        project_root = os.path.dirname(os.path.dirname(current_app.root_path))
        output_dir = os.path.join(project_root, "data", "outputs")
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir, mode=0o777, exist_ok=True)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_file = os.path.join(output_dir, f"integrated_story_{story_id}_{timestamp}.json")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump({
                'story_id': story_id,
                'integrated_at': timestamp,
                'story': story_data
            }, f, indent=2, ensure_ascii=False)
        
        logger.info(f"Integrated story {story_id} saved to: {output_file}")
        
        return jsonify({
            'success': True,
            'message': f'Story {story_id} integrated successfully',
            'storyId': story_id,
            'output_file': output_file
        })
    except Exception as e:
        logger.error(f"Error integrating story: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error integrating story: {str(e)}',
            'details': traceback.format_exc().split('\n')[-5:]
        }), 500

@api_bp.route('/integrate-all', methods=['POST'])
@login_required
def integrate_all():
    """Integrate all user stories - save to json_output (requires authentication)"""
    try:
        data = request.json
        story_ids = data.get('storyIds', [])
        stories_data = data.get('stories', [])
        
        if not stories_data:
            return jsonify({'error': 'Stories data not provided'}), 400
        
        # Save to data/outputs
        project_root = os.path.dirname(os.path.dirname(current_app.root_path))
        output_dir = os.path.join(project_root, "data", "outputs")
        
        if not os.path.exists(output_dir):
            os.makedirs(output_dir, mode=0o777, exist_ok=True)
        
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_file = os.path.join(output_dir, f"integrated_all_stories_{timestamp}.json")
        
        with open(output_file, 'w', encoding='utf-8') as f:
            json.dump({
                'integrated_at': timestamp,
                'total_stories': len(stories_data),
                'story_ids': story_ids,
                'stories': stories_data
            }, f, indent=2, ensure_ascii=False)
        
        logger.info(f"Integrated {len(stories_data)} stories saved to: {output_file}")
        
        return jsonify({
            'success': True,
            'message': f'All {len(story_ids)} stories integrated successfully',
            'storyIds': story_ids,
            'output_file': output_file
        })
    except Exception as e:
        logger.error(f"Error integrating all stories: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error integrating stories: {str(e)}',
            'details': traceback.format_exc().split('\n')[-5:]
        }), 500


@api_bp.route('/integrate-selected-github', methods=['POST'])
@login_required
def integrate_selected_github():
    """Create GitHub Issues for selected user stories (one issue per story)."""
    try:
        user: User = current_user

        if not user.github_access_token or not user.github_username:
            return jsonify({'error': 'GitHub is not connected for this user'}), 400

        if not user.github_repo:
            return jsonify({'error': 'GitHub repository is not configured'}), 400

        data = request.json or {}
        stories = data.get('stories', [])

        if not stories:
            return jsonify({'error': 'No stories provided'}), 400

        # Use github_owner if set, otherwise fallback to github_username
        owner = user.github_owner if user.github_owner else user.github_username

        if not owner:
            return jsonify({'error': 'GitHub owner/organization is not configured'}), 400

        # GitHub Issues API endpoint
        github_api_url = f"https://api.github.com/repos/{owner}/{user.github_repo}/issues"

        headers = {
            'Authorization': f"token {user.github_access_token}",
            'Accept': 'application/vnd.github+json'
        }

        created_issues = []
        errors = []

        logger.info(f"Creating {len(stories)} GitHub issues for user {user.email}")

        # Create one issue per story
        for story in stories:
            try:
                story_id = story.get('id', 'N/A')

                # Format issue title
                issue_title = story.get('title', f"User Story {story_id}")

                # Format issue body (Markdown)
                description = story.get('description', 'No description provided')
                definition_of_done = story.get('definitionOfDone', 'N/A')
                test_cases = story.get('testCases', 'N/A')

                issue_body = f"""## Description
{description}

## Definition of Done
{definition_of_done}

## Test Cases
{test_cases}

---
*Created by user-story-automation on {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC*
*User: {user.github_username}*
"""

                # Create the issue
                issue_data = {
                    'title': issue_title,
                    'body': issue_body,
                    'labels': ['user-story', 'automated']
                }

                response = requests.post(github_api_url, headers=headers, json=issue_data)

                if response.status_code == 201:
                    issue_info = response.json()
                    created_issues.append({
                        'story_id': story_id,
                        'issue_number': issue_info['number'],
                        'issue_url': issue_info['html_url']
                    })
                    logger.info(f"Created issue #{issue_info['number']} for story {story_id}")
                else:
                    try:
                        error_msg = response.json().get('message', 'Unknown error')
                    except Exception:
                        error_msg = response.text[:200]
                    errors.append(f"Story {story_id}: {error_msg}")
                    logger.error(f"Failed to create issue for story {story_id}: {response.status_code} - {error_msg}")

            except Exception as e:
                errors.append(f"Story {story.get('id', 'N/A')}: {str(e)}")
                logger.error(f"Exception creating issue for story {story.get('id')}: {e}")

        # Return results
        if created_issues:
            return jsonify({
                'success': True,
                'created_issues': created_issues,
                'total_created': len(created_issues),
                'total_failed': len(errors),
                'errors': errors if errors else None
            })
        else:
            return jsonify({
                'error': 'Failed to create any issues',
                'details': errors
            }), 400

    except Exception as e:
        logger.error(f"Error creating GitHub issues: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error creating GitHub issues: {str(e)}'
        }), 500


@api_bp.route('/github-config', methods=['GET'])
@login_required
def get_github_config():
    """Return current user's GitHub repo configuration (excluding access token)."""
    if not current_user.is_authenticated:
        return jsonify({'authenticated': False}), 401

    return jsonify({
        'authenticated': True,
        'github': {
            'username': current_user.github_username,
            'owner': current_user.github_owner,
            'repo': current_user.github_repo,
            'branch': current_user.github_branch,
            'folder': current_user.github_folder
        }
    })


@api_bp.route('/github-config', methods=['POST'])
@login_required
def update_github_config():
    """Update current user's GitHub repo configuration (repo, branch, folder)."""
    try:
        data = request.json or {}
        owner = data.get('owner')
        repo = data.get('repo')
        branch = data.get('branch') or 'main'
        folder = data.get('folder')

        if not repo:
            return jsonify({'error': 'Repository name (repo) is required'}), 400

        user: User = current_user
        if owner:
            user.github_owner = owner
        user.github_repo = repo
        user.github_branch = branch
        user.github_folder = folder
        db.session.commit()

        return jsonify({
            'success': True,
            'github': {
                'username': user.github_username,
                'owner': user.github_owner,
                'repo': user.github_repo,
                'branch': user.github_branch,
                'folder': user.github_folder
            }
        })
    except Exception as e:
        logger.error(f"Error updating GitHub config: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error updating GitHub config: {str(e)}'
        }), 500


@api_bp.route('/github-disconnect', methods=['POST'])
@login_required
def disconnect_github():
    """Disconnect GitHub by removing access token and configuration."""
    try:
        user: User = current_user
        user.github_username = None
        user.github_access_token = None
        user.github_repo = None
        user.github_branch = None
        user.github_folder = None
        db.session.commit()

        logger.info(f"GitHub disconnected for user {user.email}")
        return jsonify({'success': True})
    except Exception as e:
        logger.error(f"Error disconnecting GitHub: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error disconnecting GitHub: {str(e)}'
        }), 500

@api_bp.route('/export-json', methods=['POST'])
@login_required
def export_json():
    """Export user stories as JSON file (requires authentication)"""
    try:
        data = request.json
        stories = data.get('stories', [])
        
        if not stories:
            return jsonify({'error': 'No stories provided'}), 400
        
        # Create JSON export
        from flask import send_file
        import io
        
        export_data = {
            'exported_at': datetime.now().isoformat(),
            'user': current_user.email,
            'total_stories': len(stories),
            'stories': stories
        }
        
        # Create in-memory file
        json_str = json.dumps(export_data, indent=2, ensure_ascii=False)
        json_bytes = json_str.encode('utf-8')
        json_io = io.BytesIO(json_bytes)
        json_io.seek(0)
        
        filename = f"user_stories_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
        
        return send_file(
            json_io,
            mimetype='application/json',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        logger.error(f"Error exporting to JSON: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error exporting to JSON: {str(e)}'
        }), 500

@api_bp.route('/export-docx', methods=['POST'])
@login_required
def export_docx():
    """Export user stories as Word document (requires authentication)"""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from flask import send_file
        import io
        
        data = request.json
        stories = data.get('stories', [])
        
        if not stories:
            return jsonify({'error': 'No stories provided'}), 400
        
        # Create Word document
        doc = Document()
        
        # Add title
        title = doc.add_heading('User Stories', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        
        # Add metadata
        meta = doc.add_paragraph()
        meta.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n").bold = True
        meta.add_run(f"Total Stories: {len(stories)}\n").bold = True
        meta.add_run(f"User: {current_user.email}").bold = True
        
        doc.add_paragraph()  # Spacing
        
        # Add each story
        for idx, story in enumerate(stories, 1):
            # Story number heading
            story_heading = doc.add_heading(f"User Story {idx}", level=1)
            story_heading.runs[0].font.color.rgb = RGBColor(0, 122, 255)
            
            # Title
            if story.get('title'):
                title_para = doc.add_paragraph()
                title_para.add_run('Title: ').bold = True
                title_para.add_run(story['title'])
            
            # Description
            if story.get('description'):
                desc_para = doc.add_paragraph()
                desc_para.add_run('Description: ').bold = True
                desc_para.add_run(story['description'])
            
            # Definition of Done
            if story.get('definitionOfDone'):
                dod_para = doc.add_paragraph()
                dod_para.add_run('Definition of Done: ').bold = True

                # Split by newlines and add line breaks to preserve formatting
                dod_text = story['definitionOfDone']
                dod_lines = dod_text.split('\n')
                for i, line in enumerate(dod_lines):
                    if i > 0:
                        dod_para.add_run().add_break()  # Add line break
                    dod_para.add_run(line)

            # Test Cases
            if story.get('testCases'):
                tc_para = doc.add_paragraph()
                tc_para.add_run('Test Cases: ').bold = True

                # Split by newlines and add line breaks to preserve formatting
                tc_text = str(story['testCases'])
                tc_lines = tc_text.split('\n')
                for i, line in enumerate(tc_lines):
                    if i > 0:
                        tc_para.add_run().add_break()  # Add line break
                    tc_para.add_run(line)
            
            # Add separator between stories
            if idx < len(stories):
                doc.add_paragraph('_' * 80)
        
        # Save to in-memory file
        docx_io = io.BytesIO()
        doc.save(docx_io)
        docx_io.seek(0)
        
        filename = f"user_stories_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
        
        return send_file(
            docx_io,
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        logger.error(f"Error exporting to DOCX: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error exporting to DOCX: {str(e)}'
        }), 500


@api_bp.route('/push-to-github', methods=['POST'])
@login_required
def push_to_github():
    """Push selected user stories to GitHub repository as JSON"""
    try:
        data = request.json or {}
        stories = data.get('stories', [])

        if not stories:
            return jsonify({'error': 'No stories selected'}), 400

        user: User = current_user

        # Check if GitHub is configured with detailed validation
        logger.info(f"Push to GitHub requested by user: {user.email}")

        # Use github_owner if set, otherwise fallback to github_username
        owner = user.github_owner if user.github_owner else user.github_username

        logger.info(f"GitHub owner: {owner}, Repo: {user.github_repo}, Has token: {bool(user.github_access_token)}")

        if not owner:
            logger.error(f"GitHub owner/username missing for user {user.email}")
            return jsonify({'error': 'GitHub not connected. Please connect your GitHub account first.'}), 400

        if not user.github_access_token:
            logger.error(f"GitHub access token missing for user {user.email}")
            return jsonify({'error': 'GitHub not connected. Please connect GitHub first.'}), 400

        if not user.github_repo:
            logger.error(f"GitHub repository not configured for user {user.email}")
            return jsonify({'error': 'GitHub repository not configured. Please configure your repository settings.'}), 400

        # Prepare the JSON content
        json_content = json.dumps({
            'user_stories': stories,
            'exported_at': datetime.utcnow().isoformat(),
            'exported_by': user.github_username  # Use GitHub username instead of email
        }, indent=2)

        # Prepare file path in repo
        folder = user.github_folder.strip('/') + '/' if user.github_folder else ''
        file_path = f"{folder}user-stories-{datetime.utcnow().strftime('%Y%m%d-%H%M%S')}.json"

        # GitHub API URL
        api_url = f"https://api.github.com/repos/{owner}/{user.github_repo}/contents/{file_path}"
        logger.info(f"GitHub API URL: {api_url}")

        # Prepare the request
        headers = {
            'Authorization': f'token {user.github_access_token}',
            'Accept': 'application/vnd.github.v3+json'
        }

        # Encode content to base64
        content_encoded = base64.b64encode(json_content.encode()).decode()

        # Create the file on GitHub
        payload = {
            'message': f'Add user stories export - {len(stories)} stories',
            'content': content_encoded,
            'branch': user.github_branch or 'main'
        }

        response = requests.put(api_url, headers=headers, json=payload)

        if response.status_code in [201, 200]:
            result = response.json()
            logger.info(f"Successfully pushed {len(stories)} stories to GitHub for user {user.email}")
            return jsonify({
                'success': True,
                'message': f'Successfully pushed {len(stories)} stories to GitHub',
                'file_url': result.get('content', {}).get('html_url'),
                'file_path': file_path
            })
        else:
            logger.error(f"GitHub API error: {response.status_code} - {response.text}")
            return jsonify({
                'error': f'GitHub API error: {response.status_code}',
                'details': response.json().get('message', 'Unknown error')
            }), response.status_code

    except Exception as e:
        logger.error(f"Error pushing to GitHub: {e}")
        logger.error(traceback.format_exc())
        return jsonify({
            'error': f'Error pushing to GitHub: {str(e)}'
        }), 500
